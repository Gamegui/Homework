#!/usr/bin/env python3
"""Презентация: Ключевые военные победы России (Основы российской государственности)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x6B, 0x72, 0x80)
LIGHT = RGBColor(0xF1, 0xEC, 0xDF)

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def box(slide, l, t, w, h, fill_color=None):
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    s.line.fill.background()
    if fill_color:
        s.fill.solid()
        s.fill.fore_color.rgb = fill_color
    return s

def text_in(shape, text, size=20, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    shape.text_frame.clear()
    shape.text_frame.word_wrap = True
    p = shape.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return shape.text_frame

def bullets(shape, items, size=18, color=DARK, space=Pt(8)):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, (head, body) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = space
        p.level = 0
        if head:
            r = p.add_run()
            r.text = head
            r.font.size = Pt(size)
            r.font.bold = True
            r.font.color.rgb = NAVY
        if body:
            r = p.add_run()
            r.text = body
            r.font.size = Pt(size)
            r.font.color.rgb = color
    return tf

def footer(slide, n, total):
    t = slide.shapes.add_textbox(Inches(0.6), Inches(6.95), Inches(11), Inches(0.35))
    text_in(t, "Основы российской государственности  •  Ключевые военные победы России",
            size=11, color=GRAY)
    t2 = slide.shapes.add_textbox(Inches(12.1), Inches(6.95), Inches(0.7), Inches(0.35))
    text_in(t2, f"{n}/{total}", size=11, color=GRAY, align=PP_ALIGN.RIGHT)

def header(slide, title, kicker=None):
    box(slide, 0, 0, 13.33, 1.55, NAVY)
    box(slide, 0, 1.55, 13.33, 0.07, GOLD)
    t = slide.shapes.add_textbox(Inches(0.7), Inches(0.25), Inches(11.9), Inches(1.1))
    tf = t.text_frame
    tf.word_wrap = True
    if kicker:
        p = tf.paragraphs[0]
        p.text = kicker
        p.font.size = Pt(14)
        p.font.color.rgb = GOLD
        p.font.bold = True
        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(30)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
    else:
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(30)
        p.font.bold = True
        p.font.color.rgb = WHITE

def note(slide, text):
    slide.notes_slide.placeholders[1].text = text

TOTAL = 12

# ---------- 1. Титул ----------
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
box(s, 0.7, 1.0, 1.6, 0.08, GOLD)
t = s.shapes.add_textbox(Inches(0.7), Inches(1.4), Inches(11.9), Inches(1.4))
text_in(t, "Ключевые военные победы России", size=44, bold=True, color=WHITE)
t = s.shapes.add_textbox(Inches(0.7), Inches(2.9), Inches(11.9), Inches(0.7))
text_in(t, "От Ледового побоища до Великой Победы: как победы сформировали российскую государственность",
        size=20, color=LIGHT)
box(s, 0.7, 4.1, 11.9, 1.35, RGBColor(0x24, 0x38, 0x63))
t = s.shapes.add_textbox(Inches(1.0), Inches(4.3), Inches(11.3), Inches(1.0))
tf = t.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Предмет: Основы российской государственности  •  2026"
p.font.size = Pt(18)
p.font.color.rgb = WHITE
p2 = tf.add_paragraph()
p2.text = "Критерий отбора: победа изменила ход истории и укрепила государство, а не просто выиграла сражение"
p2.font.size = Pt(16)
p2.font.color.rgb = LIGHT
note(s, "Вступление 30 секунд: тема доклада, почему победы важны для государственности. "
        "Сразу обозначь критерий отбора — это покажет, что работа осмысленная, а не список из интернета.")

# ---------- 2. Почему это про государственность ----------
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
header(s, "Почему победы — это про государственность?", "Введение")
b = s.shapes.add_textbox(Inches(0.7), Inches(2.1), Inches(7.6), Inches(4.4))
bullets(b, [
    ("Суверенитет. ", "Каждая победа — отстоявшая независимость: от крестоносцев, Орды, поляков, шведов, французов, нацистов."),
    ("Единство. ", "Куликово поле, ополчение 1612-го, народная война 1812 и 1941 годов — победы всего народа, а не только армии."),
    ("Преемственность. ", "От Руси через империю к современной России — непрерывная линия защиты Отечества."),
    ("Ценности. ", "Служение, самопожертвование, единство — то, что курс называет традиционными ценностями."),
])
c = box(s, 8.7, 2.1, 3.9, 4.4, LIGHT)
c.text_frame.clear()
c.text_frame.word_wrap = True
p = c.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.LEFT
r = p.add_run()
r.text = "Тезис доклада\n"
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = NAVY
r = p.add_run()
r.text = "\nВоенные победы России — это вехи становления и сохранения российской государственности."
r.font.size = Pt(17)
r.font.color.rgb = DARK
c.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
c.text_frame.margin_left = Inches(0.25)
c.text_frame.margin_right = Inches(0.25)
footer(s, 2, TOTAL)
note(s, "Этот слайд — главный для предмета. Препод по Основам государственности ждёт не пересказ битв, "
        "а связь с понятиями курса: суверенитет, единство, ценности, преемственность.")

# ---------- Боевые слайды ----------
battles = [
    ("Ледовое побоище", "1242  •  Александр Невский",
     [("Противник: ", "рыцари Ливонского ордена на Чудском озере."),
      ("Что было: ", "окружение и разгром рыцарского войска; день славы — 18 апреля."),
      ("Итог: ", "остановлена экспансия крестоносцев на восток.")],
     "Для государственности: Русь сохранила северо-запад, веру и независимость Новгорода и Пскова.",
     "Первая веха: защита молодой Руси с запада. Невский — символ выбора своего пути развития."),
    ("Куликовская битва", "1380  •  Дмитрий Донской",
     [("Противник: ", "войско хана Мамая на Куликовом поле."),
      ("Что было: ", "решающий удар засадного полка; день славы — 21 сентября."),
      ("Итог: ", "первая крупная победа объединённого русского войска над Ордой.")],
     "Для государственности: разрушен миф о непобедимости Орды, начало единства русских земель вокруг Москвы.",
     "Несмотря на то что иго продлится ещё 100 лет, значение — в рождении общерусского единства."),
    ("Освобождение Москвы", "1612  •  Минин и Пожарский",
     [("Противник: ", "польский гарнизон в Кремле, войско гетмана Ходкевича."),
      ("Что было: ", "Второе ополчение разбило Ходкевича; 4 ноября — капитуляция гарнизона."),
      ("Итог: ", "конец Смутного времени, избрание Романовых в 1613 году.")],
     "Для государственности: народ спас государство, когда оно почти исчезло. Праздник — День народного единства.",
     "Ключевой пример: государства фактически не было — и его восстановил сам народ. Прямая иллюстрация единства."),
    ("Полтавская битва", "1709  •  Пётр I",
     [("Противник: ", "армия шведского короля Карла XII."),
      ("Что было: ", "разгром шведов под Полтавой; день славы — 10 июля."),
      ("Итог: ", "перелом Северной войны, выход к Балтике, статус великой державы.")],
     "Для государственности: рождение Российской империи (1721) как европейской великой державы.",
     "От обороны — к статусу империи. Полтава открыла путь к Балтике и Петербургу."),
    ("Бородино и победа 1812 года", "1812  •  М. И. Кутузов",
     [("Противник: ", "«Великая армия» Наполеона (600+ тысяч)."),
      ("Что было: ", "Бородинское сражение 7 сентября; пожар Москвы; изгнание врага к декабрю."),
      ("Итог: ", "гибель армии Наполеона, в 1814 году русские войска вошли в Париж.")],
     "Для государственности: Россия отстояла независимость и освободила Европу; рождение идеи Отечественной войны.",
     "Народный характер войны: ополчение, партизаны. День славы Бородина — 8 сентября."),
    ("Сталинград и Курск", "1942–1943",
     [("Сталинград: ", "июль 1942 — 2 февраля 1943; окружение и капитуляция армии Паулюса."),
      ("Курск: ", "июль–август 1943; крупнейшее танковое сражение; крах немецкого наступления."),
      ("Итог: ", "коренной перелом Великой Отечественной войны.")],
     "Для государственности: страна выстояла на грани уничтожения и переломила ход войны.",
     "Сталинград — символ стойкости. Подчеркнуть цену: потери с обеих сторон — более миллиона человек."),
    ("Взятие Берлина. Победа", "9 мая 1945",
     [("Что было: ", "Берлинская операция 16 апреля — 8 мая; капитуляция Германии."),
      ("Цена: ", "около 27 млн жизней советских граждан."),
      ("Итог: ", "спасение страны и мира от нацизма.")],
     "Для государственности: Победа — основа современной исторической памяти и системы дней воинской славы.",
     "Финал линии: от Чудского озера до Берлина — непрерывная защита Отечества. 9 мая — главный праздник."),
]

for i, (title, kicker, items, znamenie, speaker) in enumerate(battles, start=3):
    s = prs.slides.add_slide(BLANK)
    bg(s, WHITE)
    header(s, title, kicker)
    b = s.shapes.add_textbox(Inches(0.7), Inches(2.1), Inches(11.9), Inches(2.9))
    bullets(b, items, size=19)
    c = box(s, 0.7, 5.05, 11.9, 1.55, LIGHT)
    tf = c.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Значение для государственности: "
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r = p.add_run()
    r.text = znamenie
    r.font.size = Pt(18)
    r.font.color.rgb = DARK
    box(s, 0.7, 5.05, 0.09, 1.55, GOLD)
    footer(s, i, TOTAL)
    note(s, speaker + " Говори по 35–40 секунд на слайд.")

# ---------- 10. Что объединяет ----------
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
header(s, "Что объединяет все победы?", "Вывод  •  связь с курсом")
b = s.shapes.add_textbox(Inches(0.7), Inches(2.1), Inches(11.9), Inches(4.4))
bullets(b, [
    ("1. Оборонительный характер. ", "Россия защищала свою землю и независимость."),
    ("2. Побеждал народ. ", "Ополчение 1612-го, партизаны 1812-го, добровольцы 1941-го."),
    ("3. Единство в критический момент. ", "Разрозненные княжества, сословия, народы — единый фронт."),
    ("4. Долгий след. ", "Каждая победа определяла развитие страны на десятилетия вперёд."),
    ("5. Память как институт. ", "Дни воинской славы, 9 Мая, 4 Ноября — государство хранит память о победах."),
], size=19)
footer(s, 10, TOTAL)
note(s, "Главный выводной слайд. Пять пунктов — это готовый ответ на вопрос «в чём значение побед для государственности?».")

# ---------- 11. Дни воинской славы ----------
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
header(s, "Дни воинской славы: как государство хранит память", "Актуальность  •  ФЗ № 32-ФЗ от 13.03.1995")
b = s.shapes.add_textbox(Inches(0.7), Inches(2.1), Inches(7.6), Inches(4.4))
bullets(b, [
    ("18 апреля — ", "Ледовое побоище (1242)."),
    ("21 сентября — ", "Куликовская битва (1380)."),
    ("4 ноября — ", "освобождение Москвы (1612), День народного единства."),
    ("10 июля — ", "Полтава (1709).    8 сентября — Бородино (1812)."),
    ("2 февраля — ", "Сталинград (1943).    23 августа — Курск (1943)."),
    ("9 мая — ", "День Победы (1945)."),
], size=17)
c = box(s, 8.7, 2.1, 3.9, 4.4, LIGHT)
tf = c.text_frame
tf.clear()
tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.25)
tf.margin_right = Inches(0.25)
p = tf.paragraphs[0]
r = p.add_run()
r.text = "Почему это актуально\n"
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = NAVY
r = p.add_run()
r.text = "\nЗакон действующий, даты проверены на 2026 год. Государство закрепляет победы в календаре — это и есть политика исторической памяти."
r.font.size = Pt(16)
r.font.color.rgb = DARK
footer(s, 11, TOTAL)
note(s, "Слайд про актуальность: показать, что даты не из Википедии, а из действующего закона. "
        "Для курса это пример того, как государство работает с исторической памятью.")

# ---------- 12. Финал ----------
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
box(s, 0.7, 0.7, 1.6, 0.08, GOLD)
t = s.shapes.add_textbox(Inches(0.7), Inches(1.1), Inches(11.9), Inches(0.8))
text_in(t, "Вывод", size=36, bold=True, color=WHITE)
t = s.shapes.add_textbox(Inches(0.7), Inches(2.1), Inches(11.9), Inches(2.2))
tf = t.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Военные победы России — это вехи сохранения российской государственности: суверенитета, единства и независимости — от Чудского озера до Берлина."
p.font.size = Pt(22)
p.font.color.rgb = WHITE
t = s.shapes.add_textbox(Inches(0.7), Inches(4.3), Inches(11.9), Inches(2.2))
tf = t.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Источники: ФЗ № 32-ФЗ «О днях воинской славы»; Большая российская энциклопедия; Президентская библиотека им. Ельцина; Российская газета; материалы РВИО. Все даты сверены минимум по двум источникам."
p.font.size = Pt(15)
p.font.color.rgb = LIGHT
p2 = tf.add_paragraph()
p2.text = "Спасибо за внимание!"
p2.font.size = Pt(20)
p2.font.bold = True
p2.font.color.rgb = GOLD
note(s, "Финал 20–30 секунд: повторить тезис одним предложением, назвать источники, поблагодарить. "
        "Если спросят про цифры — честно сказать: по старым векам это оценки историков.")

prs.save("/home/user/Homework/klyuchevye-voennye-pobedy-rossii/prezentatsiya.pptx")
print("saved")
